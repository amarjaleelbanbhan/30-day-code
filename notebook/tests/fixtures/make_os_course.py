"""Generates a deterministic 7-lecture Operating Systems course used by the retrieval tests.

Deliberately includes: overlapping terminology across lectures, abbreviations stated in the slides,
speaker-note-only facts, a teacher-notes-only fact, and a textbook passage that conflicts with the lecture.
Run: python3 tests/fixtures/make_os_course.py   (requires python-pptx, reportlab)
"""
from pathlib import Path

from pptx import Presentation
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

OUT = Path(__file__).parent / "os"
OUT.mkdir(exist_ok=True)

# (title, [bullets or (level, text)], speaker notes)
LECTURES = {
    1: ("Introduction to Operating Systems", [
        ("What Is an Operating System?", ["An operating system is a program that acts as an intermediary between the user and the computer hardware",
                                          "Goals: execute user programs, make the computer convenient to use, use hardware efficiently"], None),
        ("Computer-System Organization", ["One or more CPUs and device controllers connected through a common bus to shared memory",
                                          "Device controllers inform the CPU that they have finished an operation by causing an interrupt"], None),
        ("Interrupts", ["An interrupt transfers control to the interrupt service routine through the interrupt vector",
                        "A trap (or exception) is a software-generated interrupt caused by an error or a user request"], "Remember: traps are how system calls enter the kernel. We will come back to this in Lecture 3."),
        ("Dual-Mode Operation", ["User mode and kernel mode, distinguished by a mode bit provided by the hardware",
                                 "Privileged instructions can only be executed in kernel mode"], None),
    ]),
    2: ("Operating-System Structures", [
        ("Operating-System Services", ["User interface: command-line interface (CLI), graphical user interface (GUI), touch screen",
                                       "Program execution, I/O operations, file-system manipulation, communications, error detection"], None),
        ("Monolithic Structure", ["The entire kernel runs as a single program in a single address space",
                                  "Fast, because there is little overhead in the system-call interface, but hard to extend"], "Traditional UNIX and Linux are the examples he wants us to remember."),
        ("Layered Approach", ["The operating system is divided into a number of layers; layer 0 is the hardware",
                              "Each layer only uses functions and services of lower-level layers"], None),
        ("Modules", ["Loadable kernel modules (LKMs) let the kernel link in additional services at boot or run time",
                     "Linux uses loadable kernel modules for device drivers and file systems"], None),
    ]),
    3: ("System Calls", [
        ("System Calls", ["System calls provide an interface to the services made available by an operating system",
                          "Typically written in C or C++; accessed through an application programming interface (API)"], None),
        ("API and System-Call Interface", ["The API specifies a set of functions available to programmers, e.g. the POSIX API",
                                            "The system-call interface intercepts function calls in the API and invokes the necessary system calls within the OS"], None),
        ("Parameter Passing", ["Pass parameters in registers",
                               "Store parameters in a block (table) in memory and pass the address of the block in a register",
                               "Push parameters onto the stack by the program and pop them off by the operating system"], "Block method is what Linux uses when there are more than five parameters."),
        ("Types of System Calls", ["Process control: fork(), exec(), wait(), exit()",
                                   "File management, device management, information maintenance, communications, protection"], None),
    ]),
    4: ("Kernel Architectures", [
        ("Microkernels", ["Moves as much functionality as possible from the kernel into user space",
                          "Communication between user modules uses message passing",
                          "Benefits: easier to extend, easier to port, more reliable and more secure"], "Mach is the classic microkernel example; QNX is used in cars because a crashed driver does not bring the system down."),
        ("Microkernel Performance", ["Performance overhead of user-space to kernel-space communication",
                                     "Each message must be copied between services, and the OS must switch from one process to the next"], None),
        ("Hybrid Systems", ["Most modern operating systems combine monolithic and microkernel approaches",
                            "macOS and iOS use the Darwin hybrid kernel (Mach microkernel plus BSD UNIX)"], None),
        ("Inter-Process Communication", ["Inter-process communication (IPC) lets processes exchange data",
                                         "Two models: shared memory and message passing"], None),
    ]),
    5: ("Processes", [
        ("Process Concept", ["A process is a program in execution",
                             "A program is a passive entity stored on disk; a process is an active entity",
                             "A program becomes a process when an executable file is loaded into memory"], "Program vs process is asked in the exam every year."),
        ("Process in Memory", ["Text section: the executable code",
                               "Data section: global variables",
                               "Heap: memory that is dynamically allocated during program run time",
                               "Stack: temporary data storage when invoking functions, such as function parameters, return addresses and local variables"], None),
        ("Process State", ["As a process executes, it changes state",
                           "New, Ready, Running, Waiting, Terminated: five states",
                           (1, "Only one process can be running on any processor core at any instant")], "Draw the process state diagram in the exam, with every transition labelled."),
        ("Process Control Block (PCB)", ["Each process is represented in the OS by a process control block",
                                         "Contains process state, program counter, CPU registers, CPU-scheduling information, memory-management information, accounting and I/O status"],
         "The PCB is also called task control block. He stressed this one: exam favourite."),
        ("Context Switch", ["When the CPU switches to another process, the system saves the state of the old process and loads the saved state of the new process",
                            "Context-switch time is pure overhead; the system does no useful work while switching"], None),
    ]),
    6: ("Threads", [
        ("Thread Overview", ["A thread is a basic unit of CPU utilization",
                             "It comprises a thread ID, a program counter, a register set and a stack",
                             "Threads of the same process share its code section, data section and open files"], None),
        ("Benefits of Multithreading", ["Responsiveness, resource sharing, economy, scalability",
                                        "Creating a thread is cheaper than creating a process"], None),
        ("Multicore Programming", ["Concurrency vs parallelism: a single core can provide concurrency by interleaving threads",
                                   "Challenges: dividing activities, balance, data splitting, data dependency, testing and debugging"], None),
        ("Multithreading Models", ["Many-to-one, one-to-one, many-to-many",
                                   "Linux and Windows use the one-to-one model between user threads and kernel threads"], "Many-to-one cannot run threads in parallel on multicore systems; he said this will be a short question."),
    ]),
    7: ("CPU Scheduling", [
        ("Basic Concepts", ["Maximum CPU utilization is obtained with multiprogramming",
                            "Process execution consists of a cycle of CPU execution and I/O wait: the CPU-I/O burst cycle"], None),
        ("CPU Scheduler", ["The CPU scheduler selects a process from the processes in the ready queue and allocates the CPU to it",
                           "Scheduling can be preemptive or nonpreemptive"], None),
        ("Dispatcher", ["The dispatcher gives control of the CPU to the process selected by the scheduler",
                        "Dispatch latency: the time it takes for the dispatcher to stop one process and start another"], None),
        ("Scheduling Criteria", ["CPU utilization, throughput, turnaround time, waiting time, response time"], None),
        ("First-Come, First-Served (FCFS) Scheduling", ["The process that requests the CPU first is allocated the CPU first",
                                                         "Example: P1 = 24 ms, P2 = 3 ms, P3 = 3 ms; average waiting time = (0 + 24 + 27) / 3 = 17 ms"], None),
        ("Shortest-Job-First (SJF) Scheduling", ["Associate with each process the length of its next CPU burst and pick the shortest",
                                                  "SJF gives the minimum average waiting time for a given set of processes"],
         "SJF is optimal but it cannot be implemented exactly, because the length of the next CPU burst is not known; we can only predict it."),
        ("Round Robin (RR) Scheduling", ["Each process gets a small unit of CPU time, the time quantum, usually 10 to 100 milliseconds",
                                          "After the quantum expires the process is preempted and added to the end of the ready queue"], None),
        ("Priority Scheduling", ["The CPU is allocated to the process with the highest priority",
                                 "Problem: starvation. Solution: aging, which gradually increases the priority of waiting processes"], None),
    ]),
}


def make_pptx(n, title, slides):
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[0])
    s.shapes.title.text = f"Lecture {n}: {title}"
    s.placeholders[1].text = "Operating Systems"
    for stitle, bullets, notes in slides:
        s = p.slides.add_slide(p.slide_layouts[1])
        s.shapes.title.text = stitle
        tf = s.placeholders[1].text_frame
        first = True
        for b in bullets:
            level, text = b if isinstance(b, tuple) else (0, b)
            para = tf.paragraphs[0] if first else tf.add_paragraph()
            para.text, para.level, first = text, level, False
        if notes:
            s.notes_slide.notes_text_frame.text = notes
    p.save(OUT / f"lecture{n:02d}.pptx")


for n, (title, slides) in LECTURES.items():
    make_pptx(n, title, slides)

# Textbook excerpt (reference book). Page 2 deliberately uses an older three-state model that conflicts with Lecture 5.
BOOK = [
    ("Chapter 3: Processes", ["Informally, a process is a program in execution. The status of the current activity",
                              "of a process is represented by the value of the program counter and the contents of the registers."]),
    ("3.1 Process States (simplified model)", ["In the simplified model used in this book, a process is always in one of three states:",
                                              "running, ready, or blocked. A blocked process is waiting for some event such as I/O completion."]),
    ("3.4 Interprocess Communication", ["Cooperating processes require an interprocess communication mechanism.",
                                        "Semaphores are integer variables accessed only through two atomic operations, wait() and signal()."]),
]
c = canvas.Canvas(str(OUT / "os-textbook.pdf"), pagesize=A4)
for heading, lines in BOOK:
    c.setFont("Helvetica-Bold", 16)
    c.drawString(60, 770, heading)
    c.setFont("Helvetica", 11)
    for i, line in enumerate(lines):
        c.drawString(60, 740 - i * 18, line)
    c.showPage()
c.save()

# Teacher's own notes for Lecture 7 (only place the convoy effect is mentioned).
(OUT / "scheduling-teacher-notes.md").write_text(
    "# Scheduling notes from the teacher\n\n"
    "## Convoy effect\n"
    "In FCFS all the short processes wait for one long process to get off the CPU. This convoy effect results in lower CPU and device utilization.\n\n"
    "## What to practise\n"
    "Gantt charts for FCFS, SJF and Round Robin with the same workload.\n"
)
print("wrote", sorted(p.name for p in OUT.iterdir()))
